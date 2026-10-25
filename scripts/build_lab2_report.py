from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "Lab2_Report_A4.docx"
DIAGRAMS = DOCS / "a4_diagrams"
USE_CASE_IMAGE = DIAGRAMS / "use-case-a4.png"
CLASS_OVERVIEW_IMAGE = DIAGRAMS / "class-overview-a4.png"
CLASS_DETAIL_IMAGES = [
    ("Користувачі", DIAGRAMS / "class-users-a4.png"),
    ("Каталог і кінозал", DIAGRAMS / "class-catalog-a4.png"),
    ("Замовлення та оплата", DIAGRAMS / "class-sales-a4.png"),
    ("Контролери MVC", DIAGRAMS / "class-controllers-a4.png"),
]

FONT = "Times New Roman"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F4F8FB"
GRAY = "D9D9D9"


def set_run_font(run, size: float = 14, bold: bool = False, italic: bool = False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def configure_paragraph(paragraph, *, first_line: bool = True, after: float = 6):
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_after = Pt(after)
    if first_line:
        fmt.first_line_indent = Cm(1.25)
    return paragraph


def add_body(doc: Document, text: str, *, bold_lead: str | None = None):
    p = configure_paragraph(doc.add_paragraph())
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_list_item(doc: Document, text: str, *, numbered: bool = False, level: int = 0):
    style = "List Number" if numbered else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.75 + 0.6 * level)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_compact_numbered_item(doc: Document, number: int, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(f"{number}. {text}")
    set_run_font(run, 13)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, 14, bold=True)
    return p


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = GRAY, size: str = "6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float], caption: str):
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.paragraph_format.first_line_indent = Cm(1.25)
    caption_paragraph.paragraph_format.keep_with_next = True
    caption_paragraph.paragraph_format.space_before = Pt(5)
    caption_paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(caption_paragraph.add_run(caption), 12)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, (cell, text, width) in enumerate(zip(header.cells, headers, widths_cm)):
        cell.width = Cm(width)
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        set_run_font(run, 10.5, bold=True)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, (cell, text, width) in enumerate(zip(cells, values, widths_cm)):
            cell.width = Cm(width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(text)
            set_run_font(run, 10.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, 11)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure_section(section, *, landscape: bool = False, large: bool = False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(420 if large else 297)
        section.page_height = Mm(297 if large else 210)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(25)
    section.right_margin = Mm(10 if landscape else 15)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)
    section.footer.is_linked_to_previous = False
    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.clear()
    add_page_number(footer_paragraph)


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, 12)
    return p


def add_picture(doc: Document, path: Path, max_width_mm: float, max_height_mm: float):
    from PIL import Image

    with Image.open(path) as source:
        width_px, height_px = source.size
    width_mm = min(max_width_mm, max_height_mm * width_px / height_px)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Mm(width_mm))
    return p


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.extend([r_fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_report(pr_url: str | None, student_name: str | None, group: str | None, institution: str | None, department: str | None, teacher: str | None):
    required_images = [USE_CASE_IMAGE, CLASS_OVERVIEW_IMAGE] + [path for _, path in CLASS_DETAIL_IMAGES]
    if any(not path.exists() for path in required_images):
        raise FileNotFoundError("Run scripts/draw_lab2_diagrams.py before building the report")

    doc = Document()
    section = doc.sections[0]
    configure_section(section)
    section.different_first_page_header_footer = True
    section.first_page_footer.paragraphs[0].clear()

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(14)

    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in (institution or "Заклад освіти: ____________________", department or "Кафедра: ____________________"):
        run = p.add_run(line)
        set_run_font(run, 12)
        run.add_break()

    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(14)
    run = title.add_run("Звіт до лабораторної роботи 2")
    set_run_font(run, 18, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("Проєктування вебсистеми")
    set_run_font(run, 16, bold=True)
    run.add_break()
    run2 = subtitle.add_run("Аналіз предметної галузі та побудова UML-діаграм")
    set_run_font(run2, 16, bold=True)

    topic = doc.add_paragraph()
    topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    topic.paragraph_format.space_before = Pt(18)
    topic.paragraph_format.space_after = Pt(10)
    run = topic.add_run("Тема Електронна система продажу квитків до кінотеатру")
    set_run_font(run, 15, bold=True)

    for _ in range(5):
        doc.add_paragraph()

    info = doc.add_table(rows=3, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.RIGHT
    info.autofit = False
    labels = [
        ("Виконав", student_name or "________________"),
        ("Група", group or "________________"),
        ("Перевірив", teacher or "________________"),
    ]
    for row, (label, value) in zip(info.rows, labels):
        row.cells[0].width = Cm(4.5)
        row.cells[1].width = Cm(5.0)
        row.cells[0].text = label
        row.cells[1].text = value
        for cell in row.cells:
            set_cell_margins(cell, top=40, bottom=40)
            p_cell = cell.paragraphs[0]
            p_cell.paragraph_format.space_after = Pt(0)
            for cell_run in p_cell.runs:
                set_run_font(cell_run, 12)
    tbl_pr = info._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)

    for _ in range(3):
        doc.add_paragraph()
    year = doc.add_paragraph()
    year.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(year.add_run("2026"), 12)

    doc.add_page_break()

    add_heading(doc, "1 Мета роботи")
    add_body(
        doc,
        "Провести аналіз предметної галузі електронного продажу квитків до кінотеатру, визначити функціональні вимоги через варіанти використання, побудувати діаграму варіантів використання та діаграму класів як основу для подальшої реалізації бази даних і логіки вебзастосунку на PHP та Laravel.",
    )

    add_heading(doc, "2 Завдання роботи")
    tasks = [
        "визначити призначення, функції, користувачів, сценарії, сутності та вихідні дані системи",
        "описати три різні варіанти використання у короткій, поверхневій та повній формах",
        "побудувати діаграму варіантів використання з асоціацією, узагальненням, include та extend",
        "побудувати діаграму класів із моделями й контролерами та всіма заданими типами відношень",
        "оформити архітектурну документацію і розмістити її в GitHub через окрему гілку та Pull Request",
    ]
    for item in tasks:
        add_list_item(doc, item)

    add_heading(doc, "3 Аналіз предметної галузі")
    add_heading(doc, "3.1 Призначення та функції", 2)
    add_body(
        doc,
        "Система підтримує повний цикл онлайн-продажу: від пошуку фільму до формування електронного квитка. Покупець працює з актуальним розкладом і схемою залу, а адміністратор підтримує каталог та отримує дані про продажі.",
    )
    functions = [
        "перегляд афіші, описів фільмів, вікових обмежень і тривалості",
        "перегляд розкладу за датою, фільмом і залом",
        "вибір доступних місць на схемі залу",
        "тимчасове блокування місць на час оформлення замовлення",
        "онлайн-оплата та перевірка результату платежу",
        "формування електронних квитків з унікальними кодами",
        "перегляд історії покупок і повернення квитків",
        "керування фільмами, залами, місцями та сеансами",
        "формування таблиць, звітів і графіків продажів",
    ]
    for item in functions:
        add_list_item(doc, item)

    add_heading(doc, "3.2 Користувачі", 2)
    add_body(doc, "Ролі користувачів та їхні можливості наведено в таблиці 1.")
    add_table(
        doc,
        ["Користувач", "Мета", "Можливості"],
        [
            ["Відвідувач", "Знайти фільм і сеанс", "Перегляд афіші, опису та розкладу, реєстрація або вхід"],
            ["Зареєстрований покупець", "Придбати квиток", "Вибір місць, оплата, отримання й повернення квитків"],
            ["Адміністратор", "Керувати каталогом і продажами", "Керування фільмами, залами, сеансами та звітами"],
        ],
        [3.4, 4.6, 9.0],
        "Таблиця 1 — Ролі користувачів системи",
    )

    add_heading(doc, "3.3 Архітектура MVC", 2)
    add_body(doc, "Model зберігає стан і бізнес-правила. Основні моделі: User, Customer, Administrator, Movie, CinemaHall, Seat, Screening, Order, Ticket і Payment.", bold_lead="Model")
    add_body(doc, "View відображає афішу, розклад, схему залу, форму оплати, електронний квиток та адміністративні сторінки.", bold_lead="View")
    add_body(doc, "Controller приймає HTTP-запити, перевіряє дані, викликає моделі й повертає представлення або JSON-відповідь. Передбачено CatalogController, OrderController і AdminController.", bold_lead="Controller")
    add_body(
        doc,
        "Резервування місця виконується в транзакції. Для активних квитків необхідне унікальне обмеження за сеансом і місцем. Це запобігає подвійному продажу за одночасних запитів.",
    )

    add_heading(doc, "3.4 Сценарії роботи", 2)
    scenarios = [
        "Відвідувач переглядає афішу, фільтрує список і обирає сеанс.",
        "Покупець входить до системи, вибирає місця, оплачує замовлення й отримує квиток.",
        "Адміністратор створює фільм, додає сеанс і перевіряє відсутність конфлікту розкладу.",
        "Покупець повертає невикористаний квиток, після чого система звільняє місце та ініціює повернення коштів.",
    ]
    for item in scenarios:
        add_list_item(doc, item)

    add_heading(doc, "3.5 Логічні сутності та таблиці", 2)
    add_body(doc, "Основні логічні сутності та відповідні таблиці майбутньої бази даних подано в таблиці 2.")
    add_table(
        doc,
        ["Модель", "Таблиця", "Основні поля", "Призначення"],
        [
            ["User", "users", "id, name, email, password, role, status", "Облікові записи"],
            ["Movie", "movies", "id, title, description, duration_minutes, age_rating", "Каталог фільмів"],
            ["CinemaHall", "cinema_halls", "id, name, capacity, screen_type, status", "Кінозали"],
            ["Seat", "seats", "id, hall_id, row_number, seat_number, type", "Місця залу"],
            ["Screening", "screenings", "id, movie_id, hall_id, start_at, base_price", "Розклад показів"],
            ["Order", "orders", "id, customer_id, screening_id, status, total_amount", "Оформлення покупки"],
            ["Ticket", "tickets", "id, order_id, screening_id, seat_id, code", "Квиток на місце"],
            ["Payment", "payments", "id, order_id, provider, amount, status", "Оплата та повернення"],
        ],
        [2.7, 3.3, 7.0, 5.0],
        "Таблиця 2 — Логічні сутності та таблиці бази даних",
    )

    add_heading(doc, "3.6 Вихідні дані", 2)
    outputs = [
        "таблиця доступних сеансів і місць із цінами",
        "електронний квиток із кодом, назвою фільму, датою, часом, залом, рядом і місцем",
        "таблиця продажів за вибраний період",
        "графік доходу за днями",
        "діаграма заповненості залів та популярності фільмів",
    ]
    for item in outputs:
        add_list_item(doc, item)

    add_heading(doc, "4 Варіанти використання")
    add_heading(doc, "4.1 Коротка форма Переглянути афішу та розклад", 2)
    add_body(
        doc,
        "Відвідувач відкриває афішу, обирає дату або фільм і переглядає доступні сеанси. Система показує назву фільму, вікове обмеження, тривалість, зал, час початку, базову ціну та кількість вільних місць. Відвідувач може перейти до детальної інформації або реєстрації для придбання квитка.",
    )

    add_heading(doc, "4.2 Поверхнева форма Повернути придбаний квиток", 2)
    add_body(
        doc,
        "Головний сценарій. Покупець входить до системи, відкриває власні квитки, обирає невикористаний квиток і запускає повернення. Система перевіряє строк, показує суму й запитує підтвердження. Після підтвердження система скасовує квиток, звільняє місце, надсилає запит платіжному провайдеру та повідомляє покупця.",
        bold_lead="Головний сценарій.",
    )
    add_body(
        doc,
        "Альтернативні сценарії. Якщо строк повернення минув, система відхиляє операцію. Для використаного або вже скасованого квитка повернення недоступне. Якщо платіжний провайдер тимчасово не відповідає, заявка переходить у стан очікування. Для сеансу, скасованого кінотеатром, система дозволяє повне повернення.",
        bold_lead="Альтернативні сценарії.",
    )

    add_heading(doc, "4.3 Повна форма Придбати електронний квиток", 2)
    add_body(doc, "Характеристики повного варіанта використання наведено в таблиці 3.")
    use_case_rows = [
        ["Use Case Name", "Придбати електронний квиток"],
        ["Scope", "Електронна система продажу квитків до кінотеатру"],
        ["Level", "User goal"],
        ["Primary Actor", "Зареєстрований покупець"],
        ["Stakeholders and interests", "Покупець прагне швидко придбати квиток; кінотеатр має уникнути подвійного продажу й отримати оплату; адміністратор потребує достовірних даних; платіжний провайдер потребує коректного запиту"],
        ["Preconditions", "Покупець увійшов; сеанс активний; продаж відкритий; є доступні місця"],
        ["Success guarantee", "Замовлення оплачене; місця закріплені; квитки створені; покупець отримав підтвердження; продаж потрапив до звітів"],
        ["Special Requirements", "HTTPS; транзакційне резервування; ідемпотентна оплата; цільова відповідь каталогу до 2 секунд за нормального навантаження; журналювання"],
        ["Technology and Data Variations List", "Laravel MVC, Eloquent ORM, реляційна база даних, платіжний API, HTML або PDF квиток"],
        ["Frequency of Occurrence", "Проєктна оцінка: близько 60 відсотків цільових операцій; потребує перевірки після запуску"],
        ["Miscellaneous", "Час блокування місць і правила повернення задаються конфігурацією"],
    ]
    add_table(doc, ["Розділ", "Опис"], use_case_rows, [5.3, 12.7], "Таблиця 3 — Повний опис варіанта використання")

    add_heading(doc, "4.3.1 Головний успішний сценарій", 3)
    steps = [
        "Покупець відкриває афішу.",
        "Система показує доступні фільми та дати показу.",
        "Покупець обирає фільм і потрібний сеанс.",
        "Система показує схему залу, ціни та актуальний стан місць.",
        "Покупець обирає одне або кілька вільних місць.",
        "Система атомарно блокує вибрані місця на обмежений час.",
        "Система створює чернетку замовлення та розраховує загальну суму.",
        "Покупець перевіряє замовлення і за потреби вводить промокод.",
        "Система перевіряє промокод, перераховує суму й показує остаточну вартість.",
        "Покупець підтверджує замовлення та обирає спосіб оплати.",
        "Система створює платіж і перенаправляє покупця до платіжного провайдера.",
        "Покупець підтверджує оплату.",
        "Платіжний провайдер повідомляє систему про успішну операцію.",
        "Система підтверджує замовлення та остаточно закріплює місця.",
        "Система генерує окремий електронний квиток для кожного місця.",
        "Система показує квитки в кабінеті покупця та надсилає підтвердження електронною поштою.",
    ]
    for number, item in enumerate(steps, start=1):
        add_compact_numbered_item(doc, number, item)

    add_heading(doc, "4.3.2 Розширення та альтернативні сценарії", 3)
    alternatives = [
        "3а Сеанс недоступний. Система повертає покупця до актуального розкладу.",
        "5а Місце вже зайняте. Система оновлює схему залу й пропонує інше місце.",
        "6а Не вдалося створити блокування. Система звільняє частково заблоковані місця та просить повторити вибір.",
        "9а Промокод недійсний. Система пояснює причину й дозволяє продовжити без знижки.",
        "11а Закінчився час резервування. Система скасовує чернетку та звільняє місця.",
        "13а Платіж відхилено. Система дозволяє повторити оплату, поки діє блокування.",
        "13б Немає остаточної відповіді провайдера. Замовлення переходить у стан перевірки, а повторний callback обробляється ідемпотентно.",
        "15а Не вдалося створити файл квитка. Система повторює генерацію та показує код в особистому кабінеті.",
    ]
    for number, item in enumerate(alternatives, start=1):
        add_compact_numbered_item(doc, number, item)

    doc.add_page_break()
    add_heading(doc, "5 Діаграма варіантів використання")
    add_body(
        doc,
        "Взаємодію трьох акторів із системою показано на рисунку 1. Діаграма містить асоціацію, узагальнення, include та extend. Повний вихідний код UML збережено у файлі use-case-diagram.puml.",
    )
    add_picture(doc, USE_CASE_IMAGE, 155, 190)
    add_caption(doc, "Рисунок 1 — Діаграма варіантів використання системи продажу квитків")

    doc.add_page_break()
    add_heading(doc, "6 Діаграма класів")
    add_body(doc, "Огляд десяти предметних класів і трьох контролерів подано на рисунку 2. Їхні типізовані атрибути й операції наведено на рисунках 3–6. Повний UML-код усіх зв’язків збережено у файлі class-diagram.puml.")
    add_picture(doc, CLASS_OVERVIEW_IMAGE, 155, 190)
    add_caption(doc, "Рисунок 2 — Огляд класів і типів зв’язків системи продажу квитків")
    for number, (label, image_path) in enumerate(CLASS_DETAIL_IMAGES, start=3):
        doc.add_page_break()
        add_heading(doc, f"6.{number - 2} Класи: {label}")
        add_body(doc, "На рисунку наведено типи полів, параметрів і результатів методів класів цієї групи.")
        add_picture(doc, image_path, 155, 190)
        add_caption(doc, f"Рисунок {number} — Класи: {label.lower()}")

    doc.add_page_break()
    add_heading(doc, "7 Перевірка та публікація")
    add_body(doc, "Матеріали розміщено у гілці feature/system-design репозиторію лабораторної роботи 1. Повний UML-код збережено у файлах PlantUML, а сторінки звіту створено окремим скриптом візуалізації scripts/draw_lab2_diagrams.py. Локальна перевірка scripts/validate-lab2.ps1 підтвердила три актори, тринадцять класів, шістнадцять кроків головного сценарію та вісім альтернативних сценаріїв; результат — LAB_2 VALIDATION: PASS.")
    repo_p = doc.add_paragraph()
    configure_paragraph(repo_p, first_line=False)
    set_run_font(repo_p.add_run("Архітектурний опис: "), bold=True)
    add_hyperlink(repo_p, "architecture.md у GitHub", "https://github.com/malyk-denys/php-ci-cd-lab/blob/feature/system-design/docs/architecture.md")
    commit_p = doc.add_paragraph()
    configure_paragraph(commit_p, first_line=False)
    set_run_font(commit_p.add_run("Зафіксовані зміни: "), bold=True)
    add_hyperlink(commit_p, "коміт e6a5dae", "https://github.com/malyk-denys/php-ci-cd-lab/commit/e6a5dae")
    pr_p = doc.add_paragraph()
    configure_paragraph(pr_p, first_line=False)
    set_run_font(pr_p.add_run("Pull Request: "), bold=True)
    if pr_url:
        add_hyperlink(pr_p, pr_url, pr_url)
    else:
        set_run_font(pr_p.add_run("ще не створено; підключення GitHub відхилило дію через обмеження прав інтеграції"))

    add_heading(doc, "ВИСНОВКИ")
    add_body(
        doc,
        "У роботі визначено призначення, функції, ролі, сценарії та структуру даних електронної системи продажу квитків до кінотеатру. Три форми Use Case описують основні дії користувачів, а UML-діаграми задають поведінкову й структурну модель майбутнього застосунку. Запроєктовані моделі та контролери можна безпосередньо використати як основу для Eloquent-моделей, міграцій, маршрутів і контролерів Laravel. Транзакційне блокування місць та ідемпотентна обробка платежів запобігають подвійним продажам і повторному списанню коштів.",
    )
    add_heading(doc, "ПЕРЕЛІК ДЖЕРЕЛ ПОСИЛАННЯ")
    add_body(doc, "1. Лабораторна робота 2. Проєктування вебсистеми: аналіз предметної галузі, побудова діаграм варіантів використання та класів. Методичні матеріали до дисципліни «Вебпрограмування». 22 с.")
    source_p = doc.add_paragraph()
    configure_paragraph(source_p, first_line=True)
    set_run_font(source_p.add_run("2. ДСТУ 3008:2015. Інформація та документація. Звіти у сфері науки і техніки. Структура та правила оформлювання. Київ: ДП «УкрНДНЦ», 2016. "))
    add_hyperlink(source_p, "Електронна копія", "https://sumdu.edu.ua/wp-content/uploads/2025/09/dstu_3008-2015.pdf")

    doc.core_properties.title = "Лабораторна робота 2 Електронна система продажу квитків до кінотеатру"
    doc.core_properties.subject = "Проєктування веб системи та UML діаграми"
    doc.core_properties.keywords = "PHP, Laravel, UML, Use Case, Class Diagram, cinema"
    doc.core_properties.comments = "Лабораторна робота 2"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--pr-url")
    parser.add_argument("--student-name")
    parser.add_argument("--group")
    parser.add_argument("--institution")
    parser.add_argument("--department")
    parser.add_argument("--teacher")
    args = parser.parse_args()
    build_report(args.pr_url, args.student_name, args.group, args.institution, args.department, args.teacher)
